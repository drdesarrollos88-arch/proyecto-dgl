import openpyxl
import sqlite3
import os

os.makedirs('data', exist_ok=True)
db_path = 'data/clientes.db'
conn = sqlite3.connect(db_path)
c = conn.cursor()

c.execute('''
CREATE TABLE IF NOT EXISTS clientes (
    id TEXT PRIMARY KEY,
    name TEXT NOT NULL,
    rut TEXT NOT NULL,
    comuna TEXT,
    address TEXT,
    giro TEXT,
    phone TEXT,
    payment_condition TEXT,
    email TEXT,
    contact_person TEXT,
    created_at TEXT,
    updated_at TEXT
)
''')

c.execute('CREATE INDEX IF NOT EXISTS idx_clientes_name ON clientes(name)')
c.execute('CREATE INDEX IF NOT EXISTS idx_clientes_rut ON clientes(rut)')

print('Loading Cuentas.xlsx...')
wb = openpyxl.load_workbook('Cuentas.xlsx', data_only=True)
sheet = wb['Cuentas']

c.execute('DELETE FROM clientes')

batch = []
for r in range(2, sheet.max_row + 1):
    name = sheet.cell(r, 1).value
    rut = sheet.cell(r, 2).value
    comuna = sheet.cell(r, 3).value
    address = sheet.cell(r, 4).value
    giro = sheet.cell(r, 5).value
    phone = sheet.cell(r, 6).value
    payment = sheet.cell(r, 7).value
    
    if not name and not rut:
        continue
        
    s_name = str(name).strip() if name else 'Sin Nombre'
    s_rut = str(rut).strip() if rut else ''
    s_comuna = str(comuna).strip() if comuna else ''
    s_address = str(address).strip() if address else ''
    s_giro = str(giro).strip() if giro else ''
    s_phone = str(phone).strip() if phone else ''
    s_payment = str(payment).strip() if payment else ''
    
    client_id = f'cli-{r-1}'
    batch.append((
        client_id,
        s_name,
        s_rut,
        s_comuna,
        s_address,
        s_giro,
        s_phone,
        s_payment,
        '',
        '',
        '2026-09-04T00:00:00.000Z',
        '2026-09-04T00:00:00.000Z'
    ))

c.executemany('INSERT INTO clientes VALUES (?,?,?,?,?,?,?,?,?,?,?,?)', batch)
conn.commit()

c.execute('SELECT COUNT(*) FROM clientes')
count = c.fetchone()[0]
print(f'Successfully imported {count} clients into {db_path}!')

c.execute("SELECT name, rut, comuna, phone FROM clientes WHERE name LIKE '%Nettle%' LIMIT 1")
print('Sample search result:', c.fetchone())
conn.close()

