import openpyxl
import json
import os
import subprocess

os.makedirs('data', exist_ok=True)

# Generate bcrypt hash using node
def get_bcrypt_hash(password):
    cmd = f'node -e "const bcrypt = require(\'bcryptjs\'); console.log(bcrypt.hashSync(\'{password}\', 10));"'
    result = subprocess.check_output(cmd, shell=True).decode('utf-8').strip()
    return result

print('Reading USUARIOS.xlsx...')
wb_u = openpyxl.load_workbook('USUARIOS.xlsx', data_only=True)
sheet_u = wb_u.active

users = []
for r in range(2, sheet_u.max_row + 1):
    rut = sheet_u.cell(r, 1).value
    nombre = sheet_u.cell(r, 2).value
    correo = sheet_u.cell(r, 3).value
    clave = sheet_u.cell(r, 4).value
    if correo and clave:
        users.append({
            "id": f"usr-{len(users)+1}",
            "rut": str(rut).strip() if rut else "",
            "name": str(nombre).strip() if nombre else "Usuario",
            "email": str(correo).strip().lower(),
            "passwordHash": get_bcrypt_hash(str(clave).strip()),
            "role": "admin" if len(users) == 0 else "comercial",
            "createdAt": "2026-09-04T08:00:00.000Z"
        })

print(f'Loaded {len(users)} user(s): {[u["email"] for u in users]}')

print('Reading TARIFADO 2023 r1 revPCM 27-08.xlsx...')
wb = openpyxl.load_workbook('TARIFADO 2023 r1 revPCM 27-08.xlsx', data_only=True)
sheet = wb['Tarifado 2023']

current_cat = 'ENSAYOS GENERALES'
current_subcat = ''

items = []

for r in range(2, sheet.max_row + 1):
    code = sheet.cell(r, 1).value
    desig = sheet.cell(r, 2).value
    norm = sheet.cell(r, 3).value
    weight = sheet.cell(r, 4).value
    unit = sheet.cell(r, 5).value
    val_uf = sheet.cell(r, 6).value
    sku = sheet.cell(r, 9).value
    cc = sheet.cell(r, 10).value
    
    if not desig and not code:
        continue
        
    # Check if category row
    if desig and (val_uf is None or val_uf == '' or str(val_uf).strip() == ''):
        s_desig = str(desig).strip()
        s_code = str(code).strip() if code else ""
        if s_code in ['I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII', 'IX', 'X'] or 'ENSAYOS' in s_desig.upper():
            current_cat = f"{s_code} - {s_desig}" if s_code else s_desig
            current_subcat = ""
        else:
            current_subcat = f"{s_code} {s_desig}".strip()
        continue
        
    if val_uf is not None and isinstance(val_uf, (int, float)):
        items.append({
            "id": f"item-{len(items)+1}",
            "code": str(code).strip() if code is not None else "",
            "category": current_cat,
            "subcategory": current_subcat if current_subcat else current_cat,
            "designation": str(desig).strip(),
            "norm": str(norm).strip() if norm else "",
            "minWeightKg": float(weight) if isinstance(weight, (int, float)) else (0 if weight in [None, '---', ''] else str(weight)),
            "unit": str(unit).strip() if unit else "c/u",
            "ufPrice": round(float(val_uf), 4),
            "sku": str(sku).strip() if sku else "",
            "cc": str(cc).strip() if cc else "",
            "isOfficial": True,
            "updatedAt": "2026-09-04T00:00:00.000Z",
            "updatedBy": "Sistema Maestro"
        })

db_data = {
    "version": "1.0",
    "lastUpdated": "2026-09-04T08:00:00.000Z",
    "users": users,
    "tarifario": items,
    "cotizaciones": []
}

with open('data/db.json', 'w', encoding='utf-8') as f:
    json.dump(db_data, f, ensure_ascii=False, indent=2)

print(f'Successfully initialized data/db.json with {len(users)} users and {len(items)} tarifario items!')

